from datetime import date, datetime, time, timedelta
from zipfile import BadZipFile

from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.core.exceptions import (
    SpreadsheetExtractionError,
    SpreadsheetExtractionLimitError,
)
from app.services.evidence_processing.dto.extraction_result import (
    ExtractedContentBlock,
    ExtractionResult,
    SourceReference,
)
from app.services.evidence_processing.dto.validated_file import (
    ValidatedEvidenceFile,
)
from app.services.evidence_processing.extractors.base import EvidenceExtractor


class XlsxExtractor(EvidenceExtractor):
    """Extract source-attributed spreadsheet rows without modifying evidence."""

    EXTRACTOR_NAME = "openpyxl"

    def __init__(
        self,
        max_sheets: int = 100,
        max_rows_per_sheet: int = 50_000,
        max_columns_per_sheet: int = 256,
        max_cells: int = 200_000,
        max_extracted_characters: int = 2_000_000,
    ) -> None:
        limits = (
            max_sheets,
            max_rows_per_sheet,
            max_columns_per_sheet,
            max_cells,
            max_extracted_characters,
        )

        if any(limit <= 0 for limit in limits):
            raise ValueError("Spreadsheet extraction limits must be positive")

        self.max_sheets = max_sheets
        self.max_rows_per_sheet = max_rows_per_sheet
        self.max_columns_per_sheet = max_columns_per_sheet
        self.max_cells = max_cells
        self.max_extracted_characters = max_extracted_characters

    def extract(
        self,
        evidence_file: ValidatedEvidenceFile,
    ) -> ExtractionResult:
        """Read worksheets and preserve populated cell coordinates."""

        if evidence_file.detected_file_type != "xlsx":
            raise SpreadsheetExtractionError(
                "XLSX extractor received a non-XLSX evidence file"
            )

        descriptor = evidence_file.descriptor
        blocks: list[ExtractedContentBlock] = []
        populated_cells = 0
        scanned_cells = 0
        total_characters = 0
        sheet_names: list[str] = []

        try:
            with descriptor.open() as source:
                workbook = load_workbook(
                    source,
                    read_only=True,
                    data_only=False,
                    keep_links=False,
                )

                try:
                    if len(workbook.worksheets) > self.max_sheets:
                        raise SpreadsheetExtractionLimitError(
                            "Workbook exceeds the sheet limit"
                        )

                    for sheet in workbook.worksheets:
                        sheet_names.append(sheet.title)

                        if sheet.max_row > self.max_rows_per_sheet:
                            raise SpreadsheetExtractionLimitError(
                                f"Sheet '{sheet.title}' exceeds the row limit"
                            )

                        if sheet.max_column > self.max_columns_per_sheet:
                            raise SpreadsheetExtractionLimitError(
                                f"Sheet '{sheet.title}' exceeds the column limit"
                            )

                        for row_number, row in enumerate(
                            sheet.iter_rows(
                                min_row=1,
                                max_row=sheet.max_row,
                                min_col=1,
                                max_col=sheet.max_column,
                            ),
                            start=1,
                        ):
                            scanned_cells += len(row)

                            if scanned_cells > self.max_cells:
                                raise SpreadsheetExtractionLimitError(
                                    "Workbook exceeds the cell scan limit"
                                )

                            block = self._extract_row(
                                row=row,
                                sheet_name=sheet.title,
                                row_number=row_number,
                            )

                            if block is None:
                                continue

                            populated_cells += len(block.metadata["cells"])
                            total_characters += len(block.text)

                            if total_characters > self.max_extracted_characters:
                                raise SpreadsheetExtractionLimitError(
                                    "Workbook exceeds the text limit"
                                )

                            blocks.append(block)
                finally:
                    workbook.close()

        except SpreadsheetExtractionError:
            raise
        except (
            BadZipFile,
            InvalidFileException,
            XMLSyntaxError,
            OSError,
            ValueError,
            TypeError,
            KeyError,
            RecursionError,
        ) as exc:
            raise SpreadsheetExtractionError(
                "Unable to extract spreadsheet content"
            ) from exc

        return ExtractionResult(
            evidence_id=descriptor.evidence_id,
            case_id=descriptor.case_id,
            file_type="xlsx",
            extractor_name=self.EXTRACTOR_NAME,
            blocks=tuple(blocks),
            metadata={
                "sheet_count": len(sheet_names),
                "sheet_names": tuple(sheet_names),
                "extracted_row_count": len(blocks),
                "populated_cell_count": populated_cells,
                "scanned_cell_count": scanned_cells,
                "extracted_character_count": total_characters,
            },
        )

    @classmethod
    def _extract_row(
        cls,
        row: tuple,
        sheet_name: str,
        row_number: int,
    ) -> ExtractedContentBlock | None:
        """Return one block per populated row, retaining cell addresses."""

        cells = []

        for column_number, cell in enumerate(row, start=1):
            if cell.value is None:
                continue

            cell_reference = f"{get_column_letter(column_number)}{row_number}"

            cells.append(
                {
                    "reference": cell_reference,
                    "value": cls._serialize_cell_value(cell.value),
                    "data_type": cell.data_type,
                    "is_formula": cell.data_type == "f",
                }
            )

        if not cells:
            return None

        text = "\t".join(f"{cell['reference']}: {cell['value']}" for cell in cells)

        return ExtractedContentBlock(
            text=text,
            content_type="spreadsheet_row",
            source_reference=SourceReference(
                sheet_name=sheet_name,
                row_number=row_number,
            ),
            metadata={
                "cells": tuple(cells),
            },
        )

    @staticmethod
    def _serialize_cell_value(value: object) -> str:
        """Convert a cell value to a stable, serializable representation."""

        if isinstance(value, (datetime, date, time)):
            return value.isoformat()

        if isinstance(value, timedelta):
            return str(value)

        if isinstance(value, bool):
            return str(value).lower()

        return str(value)

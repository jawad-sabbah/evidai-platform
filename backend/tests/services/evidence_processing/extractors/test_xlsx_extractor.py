import hashlib
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.core.exceptions import (
    SpreadsheetExtractionError,
    SpreadsheetExtractionLimitError,
)
from app.services.evidence_processing.dto.file_descriptor import (
    EvidenceFileDescriptor,
)
from app.services.evidence_processing.dto.validated_file import (
    ValidatedEvidenceFile,
)
from app.services.evidence_processing.extractors.xlsx_extractor import (
    XlsxExtractor,
)


def make_validated_file(path: Path) -> ValidatedEvidenceFile:
    content = path.read_bytes()
    checksum = hashlib.sha256(content).hexdigest()

    descriptor = EvidenceFileDescriptor(
        evidence_id=uuid4(),
        case_id=uuid4(),
        path=path,
        original_filename=path.name,
        file_type="xlsx",
        mime_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        file_size=len(content),
        checksum=checksum,
    )

    return ValidatedEvidenceFile(
        descriptor=descriptor,
        detected_file_type="xlsx",
        verified_size=len(content),
        verified_checksum=checksum,
    )


def test_extract_single_sheet(tmp_path):
    path = tmp_path / "transactions.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(["Transaction", "Amount"])
    sheet.append(["TX-101", 5000])
    workbook.save(path)

    result = XlsxExtractor().extract(make_validated_file(path))

    assert result.file_type == "xlsx"
    assert result.extractor_name == "openpyxl"
    assert result.metadata["sheet_count"] == 1
    assert result.metadata["sheet_names"] == ("Transactions",)
    assert len(result.blocks) == 2

    assert result.blocks[1].source_reference.sheet_name == "Transactions"
    assert result.blocks[1].source_reference.row_number == 2
    assert "B2: 5000" in result.blocks[1].text


def test_extract_multiple_sheets(tmp_path):
    path = tmp_path / "evidence.xlsx"

    workbook = Workbook()
    first_sheet = workbook.active
    first_sheet.title = "Transactions"
    first_sheet["A1"] = "TX-101"

    second_sheet = workbook.create_sheet("Witnesses")
    second_sheet["A1"] = "John Doe"

    workbook.save(path)

    result = XlsxExtractor().extract(make_validated_file(path))

    assert result.metadata["sheet_names"] == (
        "Transactions",
        "Witnesses",
    )

    assert result.blocks[0].source_reference.sheet_name == "Transactions"
    assert result.blocks[1].source_reference.sheet_name == "Witnesses"


def test_preserve_sparse_cell_coordinates(tmp_path):
    path = tmp_path / "sparse.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet["C5"] = "Evidence"
    sheet["F5"] = 500
    workbook.save(path)

    result = XlsxExtractor().extract(make_validated_file(path))

    assert len(result.blocks) == 1
    assert result.blocks[0].source_reference.row_number == 5

    cells = result.blocks[0].metadata["cells"]

    assert cells[0]["reference"] == "C5"
    assert cells[0]["value"] == "Evidence"
    assert cells[1]["reference"] == "F5"
    assert cells[1]["value"] == "500"


def test_preserve_formula_without_evaluating(tmp_path):
    path = tmp_path / "formulas.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 10
    sheet["B1"] = 20
    sheet["C1"] = "=SUM(A1:B1)"
    workbook.save(path)

    result = XlsxExtractor().extract(make_validated_file(path))

    formula = result.blocks[0].metadata["cells"][2]

    assert formula["reference"] == "C1"
    assert formula["value"] == "=SUM(A1:B1)"
    assert formula["is_formula"] is True


def test_skip_empty_rows_without_changing_coordinates(tmp_path):
    path = tmp_path / "empty_rows.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "First"
    sheet["A4"] = "Fourth"
    workbook.save(path)

    result = XlsxExtractor().extract(make_validated_file(path))

    assert len(result.blocks) == 2
    assert [block.source_reference.row_number for block in result.blocks] == [1, 4]


def test_preserve_evidence_identity(tmp_path):
    path = tmp_path / "case.xlsx"

    workbook = Workbook()
    workbook.active["A1"] = "Case evidence"
    workbook.save(path)

    evidence_file = make_validated_file(path)
    result = XlsxExtractor().extract(evidence_file)

    assert result.evidence_id == evidence_file.descriptor.evidence_id
    assert result.case_id == evidence_file.descriptor.case_id


def test_reject_invalid_workbook(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"This is not an Excel workbook")

    with pytest.raises(SpreadsheetExtractionError):
        XlsxExtractor().extract(make_validated_file(path))


def test_reject_wrong_file_type(tmp_path):
    path = tmp_path / "notes.xlsx"

    workbook = Workbook()
    workbook.save(path)

    evidence_file = make_validated_file(path)

    non_xlsx_file = ValidatedEvidenceFile(
        descriptor=evidence_file.descriptor,
        detected_file_type="txt",
        verified_size=evidence_file.verified_size,
        verified_checksum=evidence_file.verified_checksum,
    )

    with pytest.raises(SpreadsheetExtractionError):
        XlsxExtractor().extract(non_xlsx_file)


def test_enforce_sheet_limit(tmp_path):
    path = tmp_path / "many_sheets.xlsx"

    workbook = Workbook()
    workbook.create_sheet("Second")
    workbook.save(path)

    with pytest.raises(SpreadsheetExtractionLimitError):
        XlsxExtractor(max_sheets=1).extract(make_validated_file(path))


def test_enforce_row_limit(tmp_path):
    path = tmp_path / "many_rows.xlsx"

    workbook = Workbook()
    workbook.active["A10"] = "Evidence"
    workbook.save(path)

    with pytest.raises(SpreadsheetExtractionLimitError):
        XlsxExtractor(max_rows_per_sheet=5).extract(make_validated_file(path))


def test_enforce_cell_scan_limit(tmp_path):
    path = tmp_path / "many_cells.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "One"
    sheet["B1"] = "Two"
    sheet["C1"] = "Three"
    workbook.save(path)

    with pytest.raises(SpreadsheetExtractionLimitError):
        XlsxExtractor(max_cells=2).extract(make_validated_file(path))


def test_enforce_extracted_text_limit(tmp_path):
    path = tmp_path / "long_text.xlsx"

    workbook = Workbook()
    workbook.active["A1"] = "A" * 100
    workbook.save(path)

    with pytest.raises(SpreadsheetExtractionLimitError):
        XlsxExtractor(max_extracted_characters=10).extract(make_validated_file(path))


def test_original_workbook_unchanged(tmp_path):
    path = tmp_path / "original.xlsx"

    workbook = Workbook()
    workbook.active["A1"] = "Important evidence"
    workbook.save(path)

    original_bytes = path.read_bytes()

    XlsxExtractor().extract(make_validated_file(path))

    assert path.read_bytes() == original_bytes


@pytest.mark.parametrize(
    "configuration",
    [
        {"max_sheets": 0},
        {"max_rows_per_sheet": -1},
        {"max_columns_per_sheet": 0},
        {"max_cells": 0},
        {"max_extracted_characters": -10},
    ],
)
def test_reject_invalid_configuration(configuration):
    with pytest.raises(ValueError):
        XlsxExtractor(**configuration)

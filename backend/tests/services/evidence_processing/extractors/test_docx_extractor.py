from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document

from app.core.exceptions import (
    DocxExtractionError,
    DocxExtractionLimitError,
)
from app.services.evidence_processing.dto.file_descriptor import (
    EvidenceFileDescriptor,
)
from app.services.evidence_processing.dto.validated_file import (
    ValidatedEvidenceFile,
)
from app.services.evidence_processing.extractors.docx_extractor import (
    DocxTextExtractor,
)


def make_validated_file(path: Path) -> ValidatedEvidenceFile:
    descriptor = EvidenceFileDescriptor(
        evidence_id=uuid4(),
        case_id=uuid4(),
        path=path,
        original_filename=path.name,
        file_type="docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        file_size=path.stat().st_size,
        checksum=None,
    )

    return ValidatedEvidenceFile(
        descriptor=descriptor,
        detected_file_type="docx",
        verified_size=path.stat().st_size,
        verified_checksum="",
    )


def test_extract_paragraphs(tmp_path):
    path = tmp_path / "report.docx"

    document = Document()
    document.add_paragraph("First witness statement")
    document.add_paragraph("Second witness statement")
    document.save(path)

    result = DocxTextExtractor().extract(make_validated_file(path))

    assert len(result.blocks) == 2
    assert result.blocks[0].text == "First witness statement"
    assert result.blocks[1].text == "Second witness statement"
    assert result.metadata["paragraph_count"] == 2


def test_extract_headings(tmp_path):
    path = tmp_path / "report.docx"

    document = Document()
    document.add_heading("Investigation Report", level=1)
    document.add_paragraph("Evidence details")
    document.save(path)

    result = DocxTextExtractor().extract(make_validated_file(path))

    assert result.blocks[0].content_type == "heading"
    assert result.blocks[0].text == "Investigation Report"
    assert result.blocks[0].metadata["style"] == "Heading 1"
    assert result.blocks[1].content_type == "paragraph"


def test_preserve_paragraph_and_table_order(tmp_path):
    path = tmp_path / "ordered.docx"

    document = Document()
    document.add_paragraph("Before table")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Transaction"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "TX-101"
    table.cell(1, 1).text = "$5,000"

    document.add_paragraph("After table")
    document.save(path)

    result = DocxTextExtractor().extract(make_validated_file(path))

    assert [block.content_type for block in result.blocks] == [
        "paragraph",
        "table_row",
        "table_row",
        "paragraph",
    ]

    assert result.blocks[0].text == "Before table"
    assert result.blocks[1].text == "Transaction\tAmount"
    assert result.blocks[2].text == "TX-101\t$5,000"
    assert result.blocks[3].text == "After table"


def test_preserve_table_source_references(tmp_path):
    path = tmp_path / "transactions.docx"

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ID"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "101"
    table.cell(1, 1).text = "500"
    document.save(path)

    result = DocxTextExtractor().extract(make_validated_file(path))

    assert result.blocks[0].source_reference.row_number == 1
    assert result.blocks[1].source_reference.row_number == 2
    assert result.blocks[1].metadata["table_number"] == 1
    assert result.blocks[1].metadata["cells"] == ("101", "500")


def test_preserve_evidence_identity(tmp_path):
    path = tmp_path / "evidence.docx"

    document = Document()
    document.add_paragraph("Original evidence")
    document.save(path)

    evidence_file = make_validated_file(path)
    result = DocxTextExtractor().extract(evidence_file)

    assert result.evidence_id == evidence_file.descriptor.evidence_id
    assert result.case_id == evidence_file.descriptor.case_id
    assert result.file_type == "docx"
    assert result.extractor_name == "python-docx"


def test_reject_corrupted_docx(tmp_path):
    path = tmp_path / "corrupted.docx"
    path.write_bytes(b"This is not a Word document")

    with pytest.raises(DocxExtractionError):
        DocxTextExtractor().extract(make_validated_file(path))


def test_reject_non_docx_input(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Plain text", encoding="utf-8")

    evidence_file = make_validated_file(path)

    non_docx_file = ValidatedEvidenceFile(
        descriptor=evidence_file.descriptor,
        detected_file_type="txt",
        verified_size=evidence_file.verified_size,
        verified_checksum=evidence_file.verified_checksum,
    )

    with pytest.raises(DocxExtractionError):
        DocxTextExtractor().extract(non_docx_file)


def test_enforce_block_limit(tmp_path):
    path = tmp_path / "many_paragraphs.docx"

    document = Document()
    document.add_paragraph("Paragraph one")
    document.add_paragraph("Paragraph two")
    document.save(path)

    extractor = DocxTextExtractor(max_blocks=1)

    with pytest.raises(DocxExtractionLimitError):
        extractor.extract(make_validated_file(path))


def test_enforce_character_limit(tmp_path):
    path = tmp_path / "long_document.docx"

    document = Document()
    document.add_paragraph("A" * 100)
    document.save(path)

    extractor = DocxTextExtractor(max_extracted_characters=10)

    with pytest.raises(DocxExtractionLimitError):
        extractor.extract(make_validated_file(path))


def test_original_document_unchanged(tmp_path):
    path = tmp_path / "original.docx"

    document = Document()
    document.add_paragraph("Important investigation evidence")
    document.save(path)

    original_bytes = path.read_bytes()

    DocxTextExtractor().extract(make_validated_file(path))

    assert path.read_bytes() == original_bytes


@pytest.mark.parametrize(
    ("max_blocks", "max_characters"),
    [
        (0, 1000),
        (-1, 1000),
        (100, 0),
    ],
)
def test_invalid_configuration(max_blocks, max_characters):
    with pytest.raises(ValueError):
        DocxTextExtractor(
            max_blocks=max_blocks,
            max_extracted_characters=max_characters,
        )

import hashlib
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.services.evidence_processing.extraction_metadata_service import (
    ExtractionMetadataService,
)


def create_pdf_bytes() -> bytes:
    buffer = BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(50, 750, "Investigation report")
    document.showPage()
    document.drawString(50, 750, "Witness statement")
    document.showPage()
    document.save()
    return buffer.getvalue()


def create_docx_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "sample.docx"

    document = Document()
    document.add_heading("Investigation Summary", level=1)
    document.add_paragraph("Witness statement")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "TX-101"
    table.cell(0, 1).text = "5000"
    document.save(path)

    return path.read_bytes()


def create_xlsx_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "sample.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet["A1"] = "Transaction"
    sheet["B1"] = "Amount"
    sheet["A2"] = "TX-101"
    sheet["B2"] = 5000
    workbook.save(path)

    return path.read_bytes()


@pytest.mark.parametrize(
    ("filename", "content", "expected_text"),
    [
        ("notes.txt", b"Important evidence", "Important evidence"),
        ("transactions.csv", b"ID,Amount\nTX-101,5000\n", "TX-101"),
        (
            "case.json",
            b'{"transaction": "TX-101"}',
            "TX-101",
        ),
        (
            "case.xml",
            b"<case><transaction>TX-101</transaction></case>",
            "TX-101",
        ),
    ],
)
def test_structured_format_extraction_pipeline(
    filename,
    content,
    expected_text,
    evidence_builder,
    validated_snapshot,
    extractor_factory,
):
    descriptor = evidence_builder(filename, content)

    with validated_snapshot(descriptor) as evidence_file:
        extractor = extractor_factory.get_extractor(evidence_file.detected_file_type)

        extraction_result = extractor.extract(evidence_file)

        enriched_result = ExtractionMetadataService().enrich(
            evidence_file,
            extraction_result,
        )

        assert enriched_result.evidence_id == descriptor.evidence_id
        assert enriched_result.case_id == descriptor.case_id

        assert enriched_result.provenance is not None
        assert enriched_result.provenance.source_checksum_sha256 == (
            hashlib.sha256(content).hexdigest()
        )

        assert any(expected_text in block.text for block in enriched_result.blocks)

        assert [
            block.source_reference.block_number for block in enriched_result.blocks
        ] == list(range(1, len(enriched_result.blocks) + 1))

    assert descriptor.path.read_bytes() == content


def test_pdf_extraction_preserves_page_references(
    evidence_builder,
    validated_snapshot,
    extractor_factory,
):
    content = create_pdf_bytes()
    descriptor = evidence_builder("report.pdf", content)

    with validated_snapshot(descriptor) as evidence_file:
        extractor = extractor_factory.get_extractor("pdf")
        extracted = extractor.extract(evidence_file)

        result = ExtractionMetadataService().enrich(
            evidence_file,
            extracted,
        )

        assert len(result.blocks) == 2

        assert [block.source_reference.page_number for block in result.blocks] == [1, 2]

        assert "Investigation report" in result.blocks[0].text
        assert "Witness statement" in result.blocks[1].text


def test_docx_extraction_preserves_content_order(
    tmp_path,
    evidence_builder,
    validated_snapshot,
    extractor_factory,
):
    content = create_docx_bytes(tmp_path)
    descriptor = evidence_builder("report.docx", content)

    with validated_snapshot(descriptor) as evidence_file:
        extractor = extractor_factory.get_extractor("docx")
        extracted = extractor.extract(evidence_file)

        result = ExtractionMetadataService().enrich(
            evidence_file,
            extracted,
        )

        assert [block.content_type for block in result.blocks] == [
            "heading",
            "paragraph",
            "table_row",
        ]

        assert result.blocks[0].text == "Investigation Summary"
        assert result.blocks[1].text == "Witness statement"
        assert "TX-101" in result.blocks[2].text


def test_xlsx_extraction_preserves_cell_references(
    tmp_path,
    evidence_builder,
    validated_snapshot,
    extractor_factory,
):
    content = create_xlsx_bytes(tmp_path)
    descriptor = evidence_builder(
        "transactions.xlsx",
        content,
    )

    with validated_snapshot(descriptor) as evidence_file:
        extractor = extractor_factory.get_extractor("xlsx")
        extracted = extractor.extract(evidence_file)

        result = ExtractionMetadataService().enrich(
            evidence_file,
            extracted,
        )

        assert len(result.blocks) == 2
        assert result.blocks[1].source_reference.sheet_name == ("Transactions")
        assert result.blocks[1].source_reference.row_number == 2

        cells = result.blocks[1].metadata["cells"]

        assert [cell["reference"] for cell in cells] == [
            "A2",
            "B2",
        ]

        assert [cell["value"] for cell in cells] == [
            "TX-101",
            "5000",
        ]
